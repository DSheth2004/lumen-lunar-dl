"""Place paper/MANUSCRIPT.md into MDPI's official Remote Sensing Word template.

Input : paper/MANUSCRIPT.md (built by build_manuscript.py)
        paper/mdpi_template/remotesensing-template.dot (MDPI, mdpi-res.com/files/word-templates/)
Output: paper/MANUSCRIPT_MDPI.docx

The template file itself is the base document, so MDPI's page set-up, headers, footers, front-page
layout and MDPI_* styles are kept.
- Front matter (article type, title, Highlights, Abstract, Keywords) is written into the template's
  own paragraphs. Author, affiliation and academic-editor fields keep MDPI's placeholders.
- The template's example body is replaced by the manuscript body, mapped onto MDPI styles:
  headings, text, before/after-list text, bullets, numbered items, three-line tables, figures,
  captions, table footers, back matter and references.
- Equations become MDPI's two-cell equation tables (formula | number), with Office Math converted by
  pandoc from the LaTeX in the source. Inline math is converted the same way.
Requires python-docx, lxml and pypandoc (pip install pypandoc_binary).
"""
from __future__ import annotations

import copy
import io
import re
import sys
import tempfile
import zipfile
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree

HERE = Path(__file__).resolve().parent
SRC_MD = HERE / "MANUSCRIPT.md"
TEMPLATE = HERE / "mdpi_template" / "remotesensing-template.dot"
OUT = HERE / "MANUSCRIPT_MDPI.docx"

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
FULL_W, TEXT_W, TEXT_IND = 10466, 7859, 2608        # dxa: full text area, MDPI text column, its indent
INLINE = re.compile(r"(\$[^$\n]+?\$)|(\*\*.+?\*\*)|(\*[^*\s][^*]*?\*)|(`[^`]+`)")


# --------------------------------------------------------------------------------------------
# template as a document
# --------------------------------------------------------------------------------------------
def open_template() -> docx.document.Document:
    """The .dot is a Word template; retype its main part as a document so python-docx opens it."""
    buf = io.BytesIO()
    with zipfile.ZipFile(TEMPLATE) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"wordprocessingml.template.main+xml", b"wordprocessingml.document.main+xml")
            zout.writestr(item, data)
    buf.seek(0)
    return docx.Document(buf)


# --------------------------------------------------------------------------------------------
# math: LaTeX -> Office Math via pandoc
# --------------------------------------------------------------------------------------------
def convert_math(snippets: list[tuple[str, bool]]) -> list[etree._Element]:
    """[(latex, display)] -> one m:oMath element per snippet, in order."""
    import pypandoc
    if not snippets:
        return []
    md = "\n\n".join(("$$%s$$" if disp else "$%s$") % tex.strip() for tex, disp in snippets)
    with tempfile.TemporaryDirectory() as td:
        out = Path(td) / "math.docx"
        pypandoc.convert_text(md, "docx", format="markdown", outputfile=str(out))
        xml = zipfile.ZipFile(out).read("word/document.xml")
    root = etree.fromstring(xml)
    paras = root.find(qn("w:body")).findall(qn("w:p"))
    res = []
    for p in paras:
        maths = p.findall(".//{%s}oMath" % M_NS)
        if len(maths) != 1:
            raise RuntimeError("math conversion: expected 1 oMath per paragraph, got %d" % len(maths))
        res.append(maths[0])
    if len(res) != len(snippets):
        raise RuntimeError("math conversion: %d snippets -> %d paragraphs" % (len(snippets), len(res)))
    return res


# --------------------------------------------------------------------------------------------
# markdown -> blocks
# --------------------------------------------------------------------------------------------
def parse_blocks(lines: list[str]) -> list[dict]:
    blocks, i = [], 0
    while i < len(lines):
        ln = lines[i]
        if not ln.strip():
            i += 1
            continue
        if ln.startswith("#"):
            lvl = len(ln) - len(ln.lstrip("#"))
            blocks.append({"t": "h", "lvl": lvl, "text": ln[lvl:].strip()})
            i += 1
        elif ln.startswith("!["):
            m = re.match(r"!\[[^\]]*\]\(([^)]+)\)", ln)
            blocks.append({"t": "img", "path": m.group(1)})
            i += 1
        elif ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append({"t": "tbl", "rows": rows})
        elif ln.startswith("$$"):
            buf = [ln]
            while not buf[-1].rstrip().endswith("$$") or len("".join(buf).strip()) < 4:
                i += 1
                buf.append(lines[i])
            i += 1
            tex = " ".join(buf).strip()[2:-2]
            tag = re.search(r"\\tag\{([^}]*)\}", tex)
            tex = re.sub(r"\\tag\{[^}]*\}", "", tex).strip().rstrip(",.")
            blocks.append({"t": "eq", "tex": tex, "num": tag.group(1) if tag else None})
        elif ln.strip() == "---":
            blocks.append({"t": "hr"})
            i += 1
        elif re.match(r"^(\s*)(- |\d+\. )", ln):
            items = []
            while i < len(lines) and lines[i].strip():
                m = re.match(r"^(\s*)(- |\d+\. )(.*)", lines[i])
                if m:
                    items.append({"lvl": len(m.group(1)) // 2, "ordered": m.group(2) != "- ",
                                  "text": m.group(3).strip()})
                else:
                    items[-1]["text"] += " " + lines[i].strip()
                i += 1
            blocks.append({"t": "list", "items": items})
        else:
            buf = []
            while i < len(lines) and lines[i].strip() and not re.match(r"^(\s*)(- |\d+\. |\||#|!\[|\$\$)", lines[i]):
                buf.append(lines[i].strip())
                i += 1
            blocks.append({"t": "p", "text": " ".join(buf)})
    return blocks


def collect_math(blocks: list[dict]) -> list[tuple[str, bool]]:
    out = []

    def scan(text):
        for m in INLINE.finditer(text):
            if m.group(1):
                out.append((m.group(1)[1:-1], False))
            elif m.group(2):
                scan(m.group(2)[2:-2])
    for b in blocks:
        if b["t"] in ("p", "h"):
            scan(b["text"])
        elif b["t"] == "list":
            for it in b["items"]:
                scan(it["text"])
        elif b["t"] == "tbl":
            for r in b["rows"]:
                for c in r:
                    scan(c)
        elif b["t"] == "eq":
            out.append((b["tex"], True))
    return out


# --------------------------------------------------------------------------------------------
# writing
# --------------------------------------------------------------------------------------------
class Writer:
    def __init__(self, doc, maths, anchor):
        self.doc, self.maths, self.anchor = doc, iter(maths), anchor
        body = doc.element.body
        self.eq_template = next(t for t in body.iterchildren(qn("w:tbl"))
                                if re.search(r"\(1\)", "".join(t.itertext())))
        self.eq_template = copy.deepcopy(self.eq_template)

    # -- low level
    def _place(self, el):
        self.anchor.addprevious(el)

    def para(self, style, text="", align=None):
        p = self.doc.add_paragraph(style=style)
        self._place(p._p)
        if text:
            self.inline(p, text)
        if align is not None:
            p.alignment = align
        return p

    def inline(self, p, text, bold=False, italic=False):
        pos = 0
        for m in INLINE.finditer(text):
            if m.start() > pos:
                self._run(p, text[pos:m.start()], bold, italic)
            if m.group(1):
                p._p.append(copy.deepcopy(next(self.maths)))
            elif m.group(2):
                self.inline(p, m.group(2)[2:-2], True, italic)
            elif m.group(3):
                self.inline(p, m.group(3)[1:-1], bold, True)
            else:
                r = self._run(p, m.group(4)[1:-1], bold, italic)
                r.font.name = "Courier New"
            pos = m.end()
        if pos < len(text):
            self._run(p, text[pos:], bold, italic)

    @staticmethod
    def _run(p, s, bold, italic):
        r = p.add_run(s)
        r.bold = bold or None
        r.italic = italic or None
        return r

    # -- blocks
    def heading(self, lvl, text):
        style = {2: "MDPI_2.1_heading1", 3: "MDPI_2.2_heading2", 4: "MDPI_2.3_heading3"}[lvl]
        self.para(style, text)

    def equation(self, num):
        tbl = copy.deepcopy(self.eq_template)
        cells = tbl.findall(".//" + qn("w:tc"))
        for c, content in ((cells[0], "math"), (cells[1], "num")):
            p = c.find(qn("w:p"))
            for ch in list(p):
                if ch.tag != qn("w:pPr"):
                    p.remove(ch)
            if content == "math":
                p.append(copy.deepcopy(next(self.maths)))
            else:
                r = OxmlElement("w:r")
                t = OxmlElement("w:t")
                t.text = "(%s)" % num if num else ""
                r.append(t)
                p.append(r)
        self._place(tbl)

    def bullets(self, items, before_style_fix=None):
        level0_ind = None
        for it in items:
            if it["ordered"]:
                p = self.para("MDPI_3.7_itemize", it["text"])
            elif it["lvl"] == 0:
                p = self.para("MDPI_3.8_bullet", it["text"])
            else:                                    # nested: dash item, indented one step further
                p = self.para("MDPI_3.8_bullet")
                numpr = OxmlElement("w:numPr")
                for tag, val in (("w:ilvl", "0"), ("w:numId", "0")):
                    e = OxmlElement(tag)
                    e.set(qn("w:val"), val)
                    numpr.append(e)
                p._p.get_or_add_pPr().append(numpr)
                p.paragraph_format.left_indent = Cm(4.6 + 0.75 + 0.6 * it["lvl"])
                p.paragraph_format.first_line_indent = Cm(-0.45)
                self.inline(p, "–\u00a0" + it["text"])

    def table(self, rows, wide):
        ncol = max(len(r) for r in rows)
        t = self.doc.add_table(rows=len(rows), cols=ncol)
        self._place(t._tbl)
        width = FULL_W if wide else TEXT_W
        # column widths: every column fits its longest body value and its longest header word
        # (values never break mid-number; headers wrap only between words). The first column, the
        # row labels, may wrap and is capped. Spare width is shared out in proportion to need.
        clean = lambda x: re.sub(r"[*$`]", "", x)
        need = []
        for j in range(ncol):
            body_len = max([len(clean(r[j])) for r in rows[1:] if j < len(r)] or [0])
            head_word = max([len(w) for w in clean(rows[0][j]).split()] or [0]) if j < len(rows[0]) else 0
            n = max(body_len, head_word) + 1.5
            need.append(min(n, 30) if j == 0 else n)
        cw = [n * 110 + 230 for n in need]                # ~110 dxa per character (bold headers) + cell margins
        scale = width / sum(cw)
        cw = [int(c * scale) for c in cw]
        cw[-1] += width - sum(cw)
        tblPr = t._tbl.tblPr
        for tag in ("w:tblW", "w:tblInd"):
            e = OxmlElement(tag)
            e.set(qn("w:w"), str(width if tag == "w:tblW" else (0 if wide else TEXT_IND)))
            e.set(qn("w:type"), "dxa")
            tblPr.append(e)
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "bottom"):
            b = OxmlElement("w:" + side)
            for k, v in (("val", "single"), ("sz", "8"), ("space", "0"), ("color", "000000")):
                b.set(qn("w:" + k), v)
            borders.append(b)
        tblPr.append(borders)
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tblPr.append(layout)
        grid = t._tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), cw):
            gc.set(qn("w:w"), str(w))
        for i, r in enumerate(rows):
            for j in range(ncol):
                cell = t.cell(i, j)
                cell.width = docx.shared.Twips(cw[j])
                p = cell.paragraphs[0]
                p.style = self.doc.styles["MDPI_4.2_table_body"]
                txt = r[j] if j < len(r) else ""
                self.inline(p, txt, bold=(i == 0))
                if i == 0:                                   # rule under the header row
                    tcPr = cell._tc.get_or_add_tcPr()
                    tb = OxmlElement("w:tcBorders")
                    b = OxmlElement("w:bottom")
                    for k, v in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", "000000")):
                        b.set(qn("w:" + k), v)
                    tb.append(b)
                    tcPr.append(tb)

    def figure(self, path):
        p = self.para("MDPI_5.2_figure", align=WD_ALIGN_PARAGRAPH.CENTER)
        p.paragraph_format.left_indent = Cm(0)
        p.add_run().add_picture(str((HERE / path).resolve()), width=Cm(17.0))


# --------------------------------------------------------------------------------------------
def main() -> int:
    text = SRC_MD.read_text(encoding="utf-8")
    title = re.search(r"^# (.+)$", text, re.M).group(1).strip()
    hl = text[text.index("## Highlights"):text.index("## Abstract")]
    abstract = text[text.index("## Abstract") + len("## Abstract"):text.index("**Keywords:**")].strip()
    keywords = re.search(r"\*\*Keywords:\*\*\s*(.+)", text).group(1).strip()
    body_md = text[text.index("## 1. Introduction"):]
    hl_find = [b.strip()[2:].replace("\n", " ") for b in re.split(r"\n(?=- )", hl.split("**What are the implications")[0])[1:]]
    hl_impl = [b.strip()[2:].replace("\n", " ") for b in re.split(r"\n(?=- )", hl.split("**What are the implications of the main findings?**")[1])[1:]]
    if len(hl_find) != 2 or len(hl_impl) != 2:
        raise SystemExit("Highlights must have exactly 2 bullets per part (MDPI): got %d and %d" % (len(hl_find), len(hl_impl)))

    blocks = parse_blocks(body_md.splitlines())
    front_math = collect_math([{"t": "p", "text": s} for s in hl_find + hl_impl + [abstract]])
    maths = convert_math(front_math + collect_math(blocks))
    if front_math:
        raise SystemExit("math in Highlights/Abstract is not supported")

    doc = open_template()
    body = doc.element.body
    kids = list(body.iterchildren())
    paras = {i: docx.text.paragraph.Paragraph(k, doc) for i, k in enumerate(kids) if k.tag == qn("w:p")}

    def set_text(i, s, bold_prefix=None):
        p = paras[i]
        for r in list(p._p):
            if r.tag != qn("w:pPr"):
                p._p.remove(r)
        if bold_prefix:
            p.add_run(bold_prefix).bold = True
        Writer.inline(writer, p, s)

    # locate template paragraphs by text, not by position
    idx = {(" ".join(paras[i].text.split()))[:40]: i for i in paras}
    find = lambda prefix: next(i for k, i in idx.items() if k.startswith(prefix))
    i_type, i_title = find("Type of the Paper"), find("Title")
    i_hl_note = find("This is an optional section")
    i_bul = [i for i in paras if paras[i].text.strip() in ("First bullet.", "Second bullet.") and i < find("Abstract")]
    i_abs = find("A single paragraph of about 200 words")
    i_kw = find("Keywords:")
    i_line = next(i for i in paras if i > i_kw and paras[i].style.name == "MDPI_1.9_line")
    i_disc = find("Disclaimer/Publisher")

    # body: remove the template's example content between the line and the disclaimer
    writer = Writer(doc, maths, kids[i_disc])          # copies the template's equation table first
    for k in kids[i_line + 1:i_disc]:
        body.remove(k)

    set_text(i_type, "Article")
    set_text(i_title, title)
    body.remove(kids[i_hl_note])
    for i, s in zip(i_bul, hl_find + hl_impl):
        set_text(i, s)
    set_text(i_abs, " ".join(abstract.split()))
    set_text(i_kw, keywords.replace("; ", "; "), bold_prefix="Keywords: ")

    # body blocks
    prev, back = None, False
    n_tab = n_fig = n_eq = 0
    for k, b in enumerate(blocks):
        nxt = blocks[k + 1] if k + 1 < len(blocks) else None
        t = b["t"]
        if t == "hr":
            back = True
        elif t == "h":
            if b["text"] == "References":
                writer.heading(2, "References")
            else:
                writer.heading(b["lvl"], b["text"])
        elif t == "eq":
            writer.equation(b["num"])
            n_eq += 1
        elif t == "img":
            writer.figure(b["path"])
            n_fig += 1
        elif t == "tbl":
            writer.table(b["rows"], wide=max(len(r) for r in b["rows"]) >= 5)
            n_tab += 1
        elif t == "list":
            if prev and prev["t"] == "h" and prev["text"] == "References":
                for it in b["items"]:
                    writer.para("MDPI_8.1_references", re.sub(r"^\d+\.\s*", "", it["text"]))
            else:
                writer.bullets(b["items"])
        elif t == "p":
            s = b["text"]
            if re.match(r"\*\*Table \d+\.\*\*", s):
                nx_wide = nxt and nxt["t"] == "tbl" and max(len(r) for r in nxt["rows"]) >= 5
                writer.para("MDPI_4.1.1_one_table_caption" if nx_wide else "MDPI_4.1_table_caption", s)
            elif re.match(r"\*\*Figure \d+\.\*\*", s):
                writer.para("MDPI_5.1.1_one_figure_caption", s)
            elif prev and prev["t"] == "tbl" and (s.startswith("ᵃ") or s.startswith("*(")):
                writer.para("MDPI_4.3_table_footer", s)
            elif back and s == "**Abbreviations**":
                writer.heading(2, "Abbreviations")
                writer.para("MDPI_3.2_text_no_indent", "The following abbreviations are used in this manuscript:")
            elif back:
                writer.para("MDPI_6.2_back_matter", s)
            elif nxt and nxt["t"] == "list":
                writer.para("MDPI_3.5_text_before_list", s)
            elif prev and prev["t"] == "list":
                writer.para("MDPI_3.6_text_after_list", s)
            elif prev and prev["t"] == "eq":
                writer.para("MDPI_3.2_text_no_indent", s)
            else:
                writer.para("MDPI_3.1_text", s)
        prev = b
    left = sum(1 for _ in writer.maths)
    if left:
        raise SystemExit("internal error: %d converted math elements unused" % left)
    doc.core_properties.title = title
    doc.save(OUT)
    print("wrote %s: %d equations, %d tables, %d figures, %d inline+display math objects"
          % (OUT.name, n_eq, n_tab, n_fig, len(maths)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
